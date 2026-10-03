"""Document loading: DOCX, PDF, bad files, metadata, synthetic labelling."""
import pytest

from backend.ingestion.document_loader import load_document
from backend.ingestion.records import DocStatus, SourceClass
from tests.conftest import make_docx, make_pdf

SYNTHETIC_HEADER = (
    "SYNTHETIC SAMPLE FOR SOFTWARE TESTING ONLY - NOT AN OFFICIAL POLICY.\n"
)


def test_docx_paragraphs_and_headings(tmp_path):
    path = make_docx(tmp_path / "a.docx", [
        ("h1", "1. Purpose"), ("p", "Test text one."), ("bullet", "Item A"),
    ])
    doc = load_document(path)
    assert doc.status == DocStatus.OK and doc.file_type == "docx"
    lines = doc.pages[0].text.splitlines()
    assert lines == ["# 1. Purpose", "Test text one.", "- Item A"]
    assert doc.pages[0].page is None


def test_docx_table_is_extracted_in_document_order(tmp_path):
    path = make_docx(tmp_path / "t.docx", [
        ("p", "Before table."),
        ("table", [["Item", "Cost"], ["Banner", "100"]]),
        ("p", "After table."),
    ])
    lines = load_document(path).pages[0].text.splitlines()
    assert lines == ["Before table.", "| Item | Cost |", "| Banner | 100 |", "After table."]


def test_docx_merged_cells_are_not_repeated(tmp_path):
    from docx import Document
    document = Document()
    table = document.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Merged title"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    path = tmp_path / "m.docx"
    document.save(str(path))
    lines = load_document(path).pages[0].text.splitlines()
    assert lines[0] == "| Merged title |"
    assert lines[1] == "| A | B |"


def test_empty_docx_is_reported_not_crashed(tmp_path):
    path = make_docx(tmp_path / "e.docx", [])
    doc = load_document(path)
    assert doc.status == DocStatus.EMPTY and doc.error


def test_zero_byte_file_is_empty(tmp_path):
    path = tmp_path / "zero.docx"
    path.write_bytes(b"")
    assert load_document(path).status == DocStatus.EMPTY


def test_corrupt_docx_fails_gracefully(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"this is not a zip file at all")
    doc = load_document(path)
    assert doc.status == DocStatus.FAILED
    assert str(tmp_path) not in (doc.error or "")      # no file-system path leaks


def test_unsupported_and_missing_files(tmp_path):
    other = tmp_path / "notes.xlsx"
    other.write_bytes(b"x")
    assert load_document(other).status == DocStatus.UNSUPPORTED
    assert load_document(tmp_path / "missing.pdf").status == DocStatus.FAILED


def test_oversized_file_is_rejected(tmp_path):
    path = make_docx(tmp_path / "big.docx", [("p", "hello world text")])
    assert load_document(path, max_file_bytes=10).status == DocStatus.FAILED


def test_metadata_is_read_from_document_text(tmp_path):
    path = make_docx(tmp_path / "m.docx", [
        ("p", "Sample document ID: TEST-001"),
        ("p", "Version: SAMPLE 0.2 | Issuing authority: None (synthetic test content)"),
        ("p", "SYNTHETIC SAMPLE FOR SOFTWARE TESTING ONLY"),
        ("h1", "1. Body"), ("p", "Some body text here."),
    ])
    doc = load_document(path)
    assert (doc.document_id, doc.version) == ("TEST-001", "SAMPLE 0.2")
    assert doc.issuing_authority is None              # "None (...)" is never turned into an authority
    assert doc.is_synthetic and doc.source_class == SourceClass.SYNTHETIC


def test_synthetic_file_in_official_folder_is_downgraded(tmp_path):
    path = make_docx(tmp_path / "x.docx", [("p", "SYNTHETIC SAMPLE FOR SOFTWARE TESTING ONLY"), ("p", "Body text.")])
    doc = load_document(path, declared_class=SourceClass.OFFICIAL)
    assert doc.source_class == SourceClass.SYNTHETIC
    assert any("NOT official" in w for w in doc.warnings)


def test_pdf_is_read_page_by_page(tmp_path):
    path = make_pdf(tmp_path / "p.pdf", ["First page text about venues.", "Second page text about budgets."])
    doc = load_document(path)
    assert doc.status == DocStatus.OK and [p.page for p in doc.pages] == [1, 2]
    assert "venues" in doc.pages[0].text and "budgets" in doc.pages[1].text


def test_pdf_without_text_is_flagged_as_needing_ocr(tmp_path):
    path = make_pdf(tmp_path / "scan.pdf", ["", ""])
    doc = load_document(path)
    assert doc.status == DocStatus.NEEDS_OCR and doc.needs_ocr
    assert "OCR" in (doc.error or "")


def test_pdf_with_some_blank_pages_warns(tmp_path):
    path = make_pdf(tmp_path / "mixed.pdf", ["Real text on the first page of the file.", ""])
    doc = load_document(path)
    assert doc.status == DocStatus.OK and doc.ocr_pages == [2]
    assert any("OCR" in w for w in doc.warnings)


def test_corrupt_pdf_fails_gracefully(tmp_path):
    pytest.importorskip("pymupdf")
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"%PDF-1.4 garbage that is not a real pdf")
    assert load_document(path).status == DocStatus.FAILED


def test_all_provided_sample_documents_load(sample_dir):
    files = sorted(p for p in sample_dir.iterdir() if p.suffix.lower() in {".pdf", ".docx"})
    assert len(files) == 5
    for path in files:
        if path.suffix == ".pdf":
            pytest.importorskip("pymupdf")
        doc = load_document(path)
        assert doc.status == DocStatus.OK, (path.name, doc.error)
        assert doc.document_id and doc.document_id.startswith("SAMPLE-EVT-")
        assert doc.is_synthetic and doc.source_class == SourceClass.SYNTHETIC
        assert doc.issuing_authority is None
