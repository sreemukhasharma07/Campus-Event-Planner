"""Shared test helpers. All documents built here are SYNTHETIC test data."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from backend.config import PROJECT_ROOT, Settings, load_settings


@pytest.fixture
def settings() -> Settings:
    """Default settings, independent of any local .env file."""
    return load_settings(env={}, project_root=PROJECT_ROOT)


@pytest.fixture
def sample_dir() -> Path:
    return PROJECT_ROOT / "documents" / "sample_policies"


def make_docx(path: Path, items: list[tuple[str, str]]) -> Path:
    """Build a DOCX. ``items`` = [("h1"|"p"|"bullet", text) or ("table", [[cells], ...])]."""
    document = Document()
    for kind, content in items:
        if kind == "h1":
            document.add_heading(content, level=1)
        elif kind == "bullet":
            document.add_paragraph(content, style="List Bullet")
        elif kind == "table":
            rows = content
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(content)
    document.save(str(path))
    return path


def make_pdf(path: Path, pages: list[str]) -> Path:
    """Build a text PDF with PyMuPDF (tests that use it skip if PyMuPDF is missing)."""
    pymupdf = pytest.importorskip("pymupdf")
    pdf = pymupdf.open()
    for text in pages:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text, fontsize=11)
    pdf.save(str(path))
    pdf.close()
    return path
