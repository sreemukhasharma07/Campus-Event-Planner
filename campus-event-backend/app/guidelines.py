"""Loads guideline documents (pdf, docx, txt, md), chunks them and picks relevant excerpts."""
import math
import re
import threading
from dataclasses import dataclass
from pathlib import Path

from .config import settings

ALLOWED = {".pdf", ".docx", ".txt", ".md"}
STOP = set("the and for are with that this from have has will been was were not but all any can may must should shall their they them its into than then also when what which who how your you our out per etc".split())


@dataclass
class Chunk:
    source: str
    page: int | None
    text: str

    def label(self) -> str:
        return f"{self.source}, p.{self.page}" if self.page else self.source


def _tokens(text: str) -> list[str]:
    return [t for t in re.findall(r"[a-z0-9]+", text.lower()) if len(t) > 2 and t not in STOP]


def _split(text: str, limit: int = 1200) -> list[str]:
    parts, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        para = " ".join(para.split())
        if not para:
            continue
        while len(para) > limit:
            cut = para.rfind(". ", 0, limit) + 1 or limit
            parts.append(para[:cut].strip())
            para = para[cut:].strip()
        if len(cur) + len(para) + 1 > limit and cur:
            parts.append(cur)
            cur = ""
        cur = f"{cur} {para}".strip()
    if cur:
        parts.append(cur)
    return parts


def _read(path: Path) -> list[tuple[int | None, str]]:
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return [(None, path.read_text(encoding="utf-8", errors="ignore"))]
    if ext == ".pdf":
        from pypdf import PdfReader

        return [(i + 1, p.extract_text() or "") for i, p in enumerate(PdfReader(str(path)).pages)]
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return [(None, "\n\n".join(lines))]
    return []


class GuidelineIndex:
    def __init__(self):
        self._lock = threading.Lock()
        self._sig = None
        self.chunks: list[Chunk] = []
        self.docs: list[dict] = []
        self._df: dict[str, int] = {}

    def _signature(self):
        return tuple(sorted((p.name, p.stat().st_mtime_ns, p.stat().st_size) for p in self._files()))

    def _files(self):
        return [p for p in settings.guidelines_dir.iterdir() if p.is_file() and p.suffix.lower() in ALLOWED]

    def refresh(self, force: bool = False):
        with self._lock:
            sig = self._signature()
            if sig == self._sig and not force:
                return
            chunks, docs = [], []
            for path in sorted(self._files(), key=lambda p: p.name.lower()):
                status, count = "indexed", 0
                try:
                    for page, text in _read(path):
                        for piece in _split(text):
                            chunks.append(Chunk(path.name, page, piece))
                            count += 1
                    if count == 0:
                        status = "no extractable text (scanned PDF?)"
                except Exception as exc:  # corrupt or unsupported file
                    status = f"could not read ({exc.__class__.__name__})"
                docs.append({"filename": path.name, "file_type": path.suffix.lower().lstrip("."),
                             "size_kb": round(path.stat().st_size / 1024), "chunks": count, "status": status,
                             "is_sample": path.name.upper().startswith("SAMPLE")})
            df: dict[str, int] = {}
            for c in chunks:
                for t in set(_tokens(c.text)):
                    df[t] = df.get(t, 0) + 1
            self.chunks, self.docs, self._df, self._sig = chunks, docs, df, sig

    def select(self, query: str, budget: int | None = None, max_chunks: int | None = None) -> list[Chunk]:
        """All chunks if they fit the budget, otherwise the most relevant ones (kept in document order)."""
        self.refresh()
        budget = budget or settings.context_chars
        if not max_chunks and sum(len(c.text) for c in self.chunks) <= budget:
            return list(self.chunks)
        q, n = set(_tokens(query)), len(self.chunks) + 1
        scored = []
        for i, c in enumerate(self.chunks):
            toks = _tokens(c.text)
            score = sum(toks.count(t) * math.log(1 + n / (1 + self._df.get(t, 0))) for t in q) / (1 + len(toks) ** 0.3)
            scored.append((score, i))
        picked, used = [], 0
        for score, i in sorted(scored, reverse=True):
            if max_chunks and len(picked) >= max_chunks:
                break
            if used + len(self.chunks[i].text) > budget:
                continue
            picked.append(i)
            used += len(self.chunks[i].text)
        if not picked and scored:
            picked = [max(scored)[1]]
        return [self.chunks[i] for i in sorted(picked)]


def format_excerpts(chunks: list[Chunk]) -> str:
    return "\n\n".join(f"[Source: {c.label()}]\n{c.text}" for c in chunks)


index = GuidelineIndex()
